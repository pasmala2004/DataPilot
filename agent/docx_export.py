"""Render a ``ReportDoc`` as a Word document (``to_docx``) with python-docx.

Same content and order as the Markdown report: Title, Summary, Key figures (table), Analysis, Charts (as
pictures), Caveats, Sources (table). Cell texts are copied from the doc as they are (the figures were
formatted in code by ``agent/facts.py``); nothing is computed or reformatted here. The result is returned as
``bytes``: the document is built in memory and no temporary file is created.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from agent.report_doc import (
    NO_ANALYSIS_NOTE,
    NO_CHARTS_NOTE,
    NO_FIGURES_NOTE,
    NO_SOURCES_NOTE,
    PANDAS_NOTE,
    ReportDoc,
    flat,
)

PICTURE_WIDTH_INCHES = 6.0
TABLE_STYLE = "Table Grid"
MONO_FONT = "Consolas"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _table(document: "Document", header: Sequence[str], rows: Sequence[Sequence[str]], *, right_align_last: bool = False) -> None:
    """A bordered table with a bold header row; ``rows`` are plain strings."""
    table = document.add_table(rows=1, cols=len(header))
    table.style = TABLE_STYLE
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = flat(text)
        if right_align_last:
            cells[-1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if right_align_last:
        table.rows[0].cells[-1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _paragraphs(document: "Document", text: str) -> None:
    """Add text as paragraphs (a blank line starts a new one; single line breaks stay inside a paragraph)."""
    blocks = [b.strip() for b in text.replace("\r\n", "\n").split("\n\n") if b.strip()]
    for block in blocks or [""]:
        document.add_paragraph(block)


def _picture(document: "Document", path: str, alt: str) -> bool:
    """Embed the chart file as a picture (read into memory). Returns False when it cannot be read."""
    try:
        data = Path(path).read_bytes()
        document.add_picture(io.BytesIO(data), width=Inches(PICTURE_WIDTH_INCHES))
    except Exception:  # noqa: BLE001 - a missing or unreadable image must not break the whole export
        return False
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph(alt)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    return True


def to_docx(doc: ReportDoc) -> bytes:
    """Build the .docx file for ``doc`` and return its bytes."""
    document = Document()
    document.core_properties.title = doc.title
    document.core_properties.author = "DataPilot"

    document.add_heading(doc.title, level=0)
    document.add_paragraph().add_run(f"Question: {flat(doc.question)}").italic = True

    document.add_heading("Summary", level=1)
    _paragraphs(document, doc.summary)

    document.add_heading("Key figures", level=1)
    if doc.key_figures:
        _table(document, ("ID", "Figure", "Value"), [(r.ids, r.figure, r.value) for r in doc.key_figures], right_align_last=True)
    else:
        document.add_paragraph(NO_FIGURES_NOTE)

    document.add_heading("Analysis", level=1)
    if doc.analysis:
        for section in doc.analysis:
            document.add_heading(f"Task {section.task_idx + 1}", level=2)
            document.add_paragraph().add_run(flat(section.task)).italic = True
            _paragraphs(document, section.text)
    else:
        document.add_paragraph(NO_ANALYSIS_NOTE)

    document.add_heading("Charts", level=1)
    if doc.charts:
        for chart in doc.charts:
            if not _picture(document, chart.path, chart.alt):
                document.add_paragraph(f"[chart file not found: {Path(chart.path).name}]")
    else:
        document.add_paragraph(NO_CHARTS_NOTE)

    document.add_heading("Caveats", level=1)
    for caveat in doc.caveats:
        document.add_paragraph(caveat, style="List Bullet")

    document.add_heading("Sources", level=1)
    if doc.sources:
        _table(document, ("Facts", "Tool", "Arguments", "tool_call_id"),
               [(r.facts, r.tool, r.arguments, r.tool_call_id) for r in doc.sources])
    else:
        document.add_paragraph(NO_SOURCES_NOTE)
    for output in doc.pandas_outputs:
        document.add_paragraph(f"{output.tool_call_id} {PANDAS_NOTE}:")
        block = document.add_paragraph()
        run = block.add_run(output.text.rstrip("\n"))
        run.font.name, run.font.size = MONO_FONT, Pt(9)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
