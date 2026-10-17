"""agent/docx_export.py: build a .docx from a ReportDoc, read it back with python-docx. No network."""

from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pytest  # noqa: E402
from docx import Document  # noqa: E402

from agent.docx_export import to_docx  # noqa: E402
from agent.report_doc import ChartRef, ReportDoc, to_markdown  # noqa: E402
from scripted import ScriptedChatModel  # noqa: E402
from test_report import GOOD, GOOD_SUMMARY, report, run_result, writer_reply  # noqa: E402,F401  (run_result is a fixture)

NUMBER = re.compile(r"\d[\d,]*(?:\.\d+)?|<0\.001")


def real_png(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    figure = plt.figure(figsize=(2, 1.5))
    plt.bar(["a", "b"], [1, 2])
    figure.savefig(path)
    plt.close(figure)


@pytest.fixture
def built(run_result, tmp_path):
    """(report dict, parsed docx) for the standard fixture, with a real PNG behind the chart."""
    real_png(tmp_path / "outputs" / "charts" / "bar_region.png")
    result = report(run_result, ScriptedChatModel(script=[writer_reply(GOOD, GOOD_SUMMARY)]), tmp_path)
    data = to_docx(result["doc"])
    assert isinstance(data, bytes)
    return result, Document(io.BytesIO(data))


def table_rows(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def test_returns_a_valid_docx_with_the_sections_in_the_markdown_order(built):
    result, docx = built
    headings = [p.text for p in docx.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Summary", "Key figures", "Analysis", "Charts", "Caveats", "Sources"]
    assert docx.paragraphs[0].style.name == "Title" and docx.paragraphs[0].text == result["doc"].title
    assert [p.text for p in docx.paragraphs if p.style.name == "Heading 2"] == ["Task 1", "Task 2", "Task 3"]
    assert docx.core_properties.title == result["doc"].title


def test_every_number_in_the_key_figures_table_equals_the_fact_text(built):
    result, docx = built
    facts = {f["fact_id"]: f["text"] for f in result["facts"]}
    header, *rows = table_rows(docx.tables[0])
    assert header == ["ID", "Figure", "Value"]
    assert len(rows) == len(result["doc"].key_figures) == 6
    for ids, figure, value in rows:
        assert value == " / ".join(facts[i.strip()] for i in ids.split(","))     # the cell IS the fact text(s)
        assert figure
    shown = {text for _, _, value in rows for text in value.split(" / ")}
    assert shown == set(facts.values())                                            # and every fact is shown once
    assert {"1,470.25", "80.00", "0.835", "0.005100"} <= shown


def test_docx_and_markdown_show_exactly_the_same_numbers(built):
    result, docx = built
    markdown = to_markdown(result["doc"])
    paragraphs = "\n".join(p.text for p in docx.paragraphs)
    cells = "\n".join(c for r in table_rows(docx.tables[0]) for c in r)                 # the Key figures table
    clean = lambda text: re.sub(r"Task \d+|\bF\d+\b", "", text)                        # noqa: E731 - drop ids and task numbers
    docx_numbers = sorted(NUMBER.findall(clean(paragraphs.split("Caveats")[0] + "\n" + cells)))
    md_numbers = sorted(NUMBER.findall(clean(markdown.split("## Caveats")[0])))
    assert docx_numbers == md_numbers


def test_picture_count_equals_chart_count(built, run_result, tmp_path):
    result, docx = built
    assert len(result["doc"].charts) == 1 and len(docx.inline_shapes) == 1
    # two charts -> two pictures
    real_png(tmp_path / "outputs" / "charts" / "second.png")
    doc = result["doc"]
    two = ReportDoc(**{**doc.__dict__, "charts": [*doc.charts, ChartRef(str(tmp_path / "outputs" / "charts" / "second.png"),
                                                                      "../charts/second.png", "second chart")]})
    assert len(Document(io.BytesIO(to_docx(two))).inline_shapes) == 2


def test_a_missing_chart_file_becomes_a_note_not_a_crash(built):
    result, _ = built
    doc = result["doc"]
    gone = ReportDoc(**{**doc.__dict__, "charts": [ChartRef("/no/such/chart.png", "x.png", "gone")]})
    docx = Document(io.BytesIO(to_docx(gone)))
    assert len(docx.inline_shapes) == 0
    assert any("chart file not found: chart.png" in p.text for p in docx.paragraphs)


def test_caveats_and_sources_are_present(built):
    result, docx = built
    bullets = [p.text for p in docx.paragraphs if p.style.name == "List Bullet"]
    assert bullets == result["doc"].caveats and len(bullets) >= 3
    assert any("Task 4 failed" in b for b in bullets) and any("computed by tools" in b for b in bullets)
    sources = table_rows(docx.tables[1])
    assert sources[0] == ["Facts", "Tool", "Arguments", "tool_call_id"]
    assert [r[3] for r in sources[1:]] == ["call_a", "call_b", "call_p", "call_c"]
    assert sources[1][0] == "F1-F4" and sources[1][1] == "aggregate"
    assert any("call_p" in p.text and "free text" in p.text for p in docx.paragraphs)       # run_pandas output block


def test_no_results_report_exports_with_notes_instead_of_empty_sections(run_result, tmp_path):
    for status in run_result["task_status"]:
        status["status"] = "failed"
    result = report(run_result, ScriptedChatModel(script=[]), tmp_path)
    docx = Document(io.BytesIO(to_docx(result["doc"])))
    text = "\n".join(p.text for p in docx.paragraphs)
    for note in ("No verified figures were produced.", "No analysis could be verified.", "Nothing could be verified"):
        assert note in text
    assert docx.tables == [] and len(docx.inline_shapes) == 0


def test_export_creates_no_files_and_needs_no_temp_files(built, monkeypatch, tmp_path):
    result, _ = built
    before = sorted(p.as_posix() for p in tmp_path.rglob("*"))

    def refuse(*a, **k):
        raise AssertionError("to_docx must not create temporary files")

    monkeypatch.setattr(tempfile, "NamedTemporaryFile", refuse)
    monkeypatch.setattr(tempfile, "mkstemp", refuse)
    monkeypatch.setattr(tempfile, "mkdtemp", refuse)
    assert to_docx(result["doc"])[:2] == b"PK"                                        # a zip container
    assert sorted(p.as_posix() for p in tmp_path.rglob("*")) == before


def test_table_cells_are_not_cut_or_escaped_like_markdown(run_result, tmp_path):
    doc = ReportDoc(title="t", question="q", summary="s", caveats=["c"],
                    key_figures=[], analysis=[], charts=[], sources=[])
    long_label = "x" * 300 + " | pipe"
    from agent.report_doc import KeyFigureRow
    doc.key_figures = [KeyFigureRow("F1", long_label, "1.00")]
    cells = table_rows(Document(io.BytesIO(to_docx(doc))).tables[0])[1]
    assert cells[1] == long_label and "\\|" not in cells[1]


def test_long_run_pandas_output_is_cut_in_the_docx_sources_too(run_result, tmp_path):
    run_result["results"][2]["output"] = {"output": "".join(f"row {i}\n" for i in range(1, 21))}
    real_png(tmp_path / "outputs" / "charts" / "bar_region.png")
    result = report(run_result, ScriptedChatModel(script=[writer_reply(GOOD, GOOD_SUMMARY)]), tmp_path)
    docx = Document(io.BytesIO(to_docx(result["doc"])))
    block = next(p.text for p in docx.paragraphs if p.text.startswith("row 1"))
    assert block == "row 1\nrow 2\nrow 3\nrow 4\nrow 5\n(truncated)"
    assert not any("row 6" in p.text for p in docx.paragraphs)
